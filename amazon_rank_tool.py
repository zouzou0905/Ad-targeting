from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl


APP_DIR = Path(__file__).resolve().parent
DEFAULT_SOURCE = Path.home() / "Desktop" / "贺卡广告关键词位置情况登记.xlsx"
ASIN_PATTERN = re.compile(r"^[A-Z0-9]{10}$")
FIXED_SCAN_PAGES = 3


def configure_console() -> None:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")


def normalize_keyword(value: str) -> str:
    return " ".join(value.strip().lower().split())


def locate_source_sheet(workbook: openpyxl.Workbook) -> tuple[Any, int, dict[str, int]]:
    required = {"关键词", "ASIN"}
    optional = {"广告活动", "竞价", "图片"}
    for worksheet in workbook.worksheets:
        for row in range(1, min(worksheet.max_row, 12) + 1):
            headers: dict[str, int] = {}
            for column in range(1, min(worksheet.max_column, 20) + 1):
                value = worksheet.cell(row, column).value
                if value is None:
                    continue
                label = str(value).strip()
                if label in required or label in optional:
                    headers[label] = column
            if required.issubset(headers):
                return worksheet, row, headers
    raise ValueError("没有找到同时包含“关键词”和“ASIN”的表头。")


def extract_tasks(source: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(source, read_only=False, data_only=False)
    worksheet, header_row, headers = locate_source_sheet(workbook)

    keyword_col = headers["关键词"]
    asin_col = headers["ASIN"]
    campaign_col = headers.get("广告活动")
    bid_col = headers.get("竞价")

    first_data_row = None
    for row in range(header_row + 1, min(worksheet.max_row, header_row + 50) + 1):
        value = worksheet.cell(row, asin_col).value
        asin = str(value).strip().upper() if value is not None else ""
        if ASIN_PATTERN.fullmatch(asin):
            first_data_row = row
            break
    if first_data_row is None:
        raise ValueError("表格中没有识别到 10 位 ASIN。")

    last_data_row = max(
        (
            row
            for row in range(first_data_row, worksheet.max_row + 1)
            if ASIN_PATTERN.fullmatch(
                str(worksheet.cell(row, asin_col).value or "").strip().upper()
            )
        ),
        default=first_data_row,
    )

    current_keyword = ""
    current_campaign = ""
    source_rows: list[dict[str, Any]] = []
    tasks_by_key: OrderedDict[str, dict[str, Any]] = OrderedDict()

    for row in range(first_data_row, last_data_row + 1):
        keyword_value = worksheet.cell(row, keyword_col).value
        campaign_value = worksheet.cell(row, campaign_col).value if campaign_col else None
        asin_value = worksheet.cell(row, asin_col).value

        if keyword_value not in (None, ""):
            current_keyword = str(keyword_value).strip()
        if campaign_value not in (None, ""):
            current_campaign = str(campaign_value).strip()

        asin = str(asin_value or "").strip().upper()
        if not ASIN_PATTERN.fullmatch(asin):
            continue
        if not current_keyword:
            raise ValueError(f"第 {row} 行 ASIN {asin} 没有对应关键词。")

        keyword_key = normalize_keyword(current_keyword)
        task = tasks_by_key.setdefault(
            keyword_key,
            {"keyword": current_keyword, "asins": []},
        )
        if asin not in task["asins"]:
            task["asins"].append(asin)

        source_rows.append(
            {
                "source_sheet": worksheet.title,
                "source_row": row,
                "campaign": current_campaign,
                "keyword": current_keyword,
                "keyword_key": keyword_key,
                "asin": asin,
                "bid": worksheet.cell(row, bid_col).value if bid_col else None,
            }
        )

    workbook.close()
    return {
        "sheet": worksheet.title,
        "source_rows": source_rows,
        "tasks": list(tasks_by_key.values()),
    }


def find_runtime_executable(relative_parts: tuple[str, ...], fallbacks: list[str]) -> str:
    runtime_root = (
        Path.home()
        / ".cache"
        / "codex-runtimes"
        / "codex-primary-runtime"
        / "dependencies"
    )
    candidate = runtime_root.joinpath(*relative_parts)
    if candidate.exists():
        return str(candidate)
    for fallback in fallbacks:
        resolved = shutil.which(fallback)
        if resolved:
            return resolved
    raise FileNotFoundError(f"找不到运行程序：{relative_parts[-1]}")


def find_chrome() -> str:
    candidates = [
        Path(os.environ.get("PROGRAMFILES", "C:/Program Files"))
        / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("PROGRAMFILES(X86)", "C:/Program Files (x86)"))
        / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("LOCALAPPDATA", "")) / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("PROGRAMFILES(X86)", "C:/Program Files (x86)"))
        / "Microsoft/Edge/Application/msedge.exe",
        Path(os.environ.get("PROGRAMFILES", "C:/Program Files"))
        / "Microsoft/Edge/Application/msedge.exe",
    ]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    raise FileNotFoundError("没有找到 Chrome 或 Edge。")


def prompt_value(label: str, default: str) -> str:
    value = input(f"{label}（直接回车使用 {default}）：").strip()
    return value or default


def occurrence_text(occurrences: list[dict[str, Any]]) -> str:
    if not occurrences:
        return "无"
    return "，".join(
        f"{item['typePosition']}（第{item['page']}页/页内{item['pagePosition']}/总位置{item['overallPosition']}）"
        for item in occurrences
    )


def module_occurrence_text(occurrences: list[dict[str, Any]]) -> str:
    if not occurrences:
        return "无"
    parts = []
    for item in occurrences:
        cards_before = item.get("cardsBefore") or 0
        where = f"第{cards_before}个搜索结果之后" if cards_before > 0 else "搜索结果最前"
        parts.append(f"{item.get('kind', '广告模块')}（第{item['page']}页/{where}）")
    return "，".join(parts)


def target_status(target: dict[str, Any]) -> str:
    parts = []
    if target.get("adOccurrences"):
        parts.append("广告位")
    if target.get("organicOccurrences"):
        parts.append("自然位")
    if target.get("adModuleOccurrences"):
        parts.append("网格外广告位")
    return "+".join(parts) if parts else "未找到"


def build_text_log(scan_payload: dict[str, Any]) -> str:
    lines = [
        "Amazon 关键词 ASIN 位置检测日志",
        f"检测时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"Amazon 站点：{scan_payload.get('baseUrl', '')}",
        f"每个关键词最多扫描：{scan_payload.get('maxPages', '')} 页",
        "",
    ]
    for index, result in enumerate(scan_payload.get("results", []), start=1):
        lines.append(f"[{index}] 关键词：{result['keyword']}")
        if result.get("error"):
            lines.append(f"扫描错误：{result['error']}")
        for summary in result.get("pageSummaries", []):
            load_state = "已稳定" if summary.get("stabilized") else "未完全稳定"
            lines.append(
                f"第 {summary['page']} 页：计入 {summary['productCards']} 个商品位置，"
                f"广告 {summary['sponsoredCards']}，自然 {summary['organicCards']}；"
                f"网格外广告模块 {summary.get('adModules', 0)} 个；"
                f"标准结果 {summary.get('standardCards', '未知')}，"
                f"可见唯一 ASIN {summary.get('uniqueDataAsins', '未知')}，"
                f"商品链接唯一 ASIN {summary.get('uniqueProductLinkAsins', '未知')}，"
                f"加载{load_state}"
            )
        for target in result.get("targets", {}).values():
            ads = target.get("adOccurrences", [])
            organic = target.get("organicOccurrences", [])
            modules = target.get("adModuleOccurrences", [])
            lines.append(
                f"ASIN {target['asin']}｜{target_status(target)}｜广告排名：{occurrence_text(ads)}｜"
                f"自然排名：{occurrence_text(organic)}"
            )
            if modules:
                lines.append(f"网格外广告位：{module_occurrence_text(modules)}")
            occurrences = sorted(ads + organic, key=lambda item: item["overallPosition"])
            shown_title = (occurrences[0].get("title") if occurrences else "") or (
                modules[0].get("title") if modules else ""
            )
            shown_url = (occurrences[0].get("url") if occurrences else "") or (
                modules[0].get("url") if modules else ""
            )
            if shown_title or shown_url:
                lines.append(f"标题：{shown_title or '（未读取到标题）'}")
                lines.append(f"链接：{shown_url or '（未读取到链接）'}")
            if ads:
                reasons = list(dict.fromkeys(item.get("sponsoredReason", "") for item in ads))
                lines.append(f"广告识别依据：{'，'.join(reason for reason in reasons if reason)}")
        lines.append("")

    lines.extend(
        [
            "========== 有位置的关键词和 ASIN 汇总 ==========",
        ]
    )
    found_count = 0
    for result in scan_payload.get("results", []):
        for target in result.get("targets", {}).values():
            if not target.get("found"):
                continue
            found_count += 1
            lines.append(
                f"关键词：{result['keyword']}｜ASIN：{target['asin']}｜"
                f"广告排名：{occurrence_text(target.get('adOccurrences', []))}｜"
                f"自然排名：{occurrence_text(target.get('organicOccurrences', []))}｜"
                f"网格外广告位：{module_occurrence_text(target.get('adModuleOccurrences', []))}"
            )
    if found_count == 0:
        lines.append("本次扫描未发现有位置的关键词和 ASIN。")
    lines.append(f"汇总：共 {found_count} 组有位置的关键词/ASIN。")
    lines.append("================================================")
    return "\n".join(lines)


def run_tool(args: argparse.Namespace) -> tuple[Path, Path]:
    source = Path(args.input).expanduser().resolve()
    if not source.exists():
        raise FileNotFoundError(f"找不到输入文件：{source}")

    print("\n正在读取表格并整理关键词与 ASIN……")
    extracted = extract_tasks(source)
    print(
        f"识别到 {len(extracted['tasks'])} 个唯一关键词、"
        f"{len(extracted['source_rows'])} 行 ASIN。"
    )
    if args.limit > 0:
        extracted["tasks"] = extracted["tasks"][: args.limit]
        print(f"本次为核对模式，只扫描前 {len(extracted['tasks'])} 个唯一关键词。")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_dir = Path(args.log_dir).expanduser().resolve() if args.log_dir else APP_DIR / "logs"
    log_dir.mkdir(parents=True, exist_ok=True)
    text_log = log_dir / f"检测日志_{timestamp}.txt"
    json_log = log_dir / f"检测日志_{timestamp}.json"

    node = find_runtime_executable(("node", "bin", "node.exe"), ["node.exe", "node"])
    chrome = find_chrome()
    with tempfile.TemporaryDirectory(prefix="amazon-rank-") as temp_dir_string:
        temp_dir = Path(temp_dir_string)
        # Debug/headless runs use an isolated profile so they never collide with
        # an interactive scan that is waiting for the user in another window.
        profile_dir = (
            temp_dir / "browser-profile"
            if args.headless
            else APP_DIR / ".runtime" / "amazon-browser-profile"
        )
        task_file = temp_dir / "tasks.json"
        result_file = temp_dir / "results.json"
        task_payload = {
            "baseUrl": args.amazon,
            "maxPages": FIXED_SCAN_PAGES,
            "delaySeconds": args.delay,
            "profileDir": str(profile_dir),
            "chromeExecutable": chrome,
            "headless": args.headless,
            "tasks": extracted["tasks"],
        }
        task_file.write_text(json.dumps(task_payload, ensure_ascii=False, indent=2), encoding="utf-8")

        command = [
            node,
            str(APP_DIR / "amazon_scanner.mjs"),
            "--input",
            str(task_file),
            "--output",
            str(result_file),
        ]
        if args.headless:
            command.extend(["--headless", "--no-setup"])

        completed = subprocess.run(command, cwd=APP_DIR, check=False)
        if completed.returncode != 0:
            raise RuntimeError("浏览器扫描未完成，未生成结果文件。")
        scan_payload = json.loads(result_file.read_text(encoding="utf-8"))

    text_log.write_text(build_text_log(scan_payload), encoding="utf-8-sig")
    json_log.write_text(json.dumps(scan_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print("\n检测完成。没有生成或修改任何 Excel 文件。")
    print(f"可读日志：{text_log}")
    print(f"原始数据日志：{json_log}")
    return text_log, json_log


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="读取 Excel 关键词与 ASIN，检测 Amazon 搜索位置。")
    parser.add_argument("--input", help="输入 Excel 文件")
    parser.add_argument("--amazon", default="https://www.amazon.com", help="Amazon 站点")
    parser.add_argument("--delay", type=float, default=4.0, help="关键词之间的等待秒数")
    parser.add_argument("--limit", type=int, default=0, help="只扫描前 N 个唯一关键词，0 表示全部")
    parser.add_argument("--log-dir", help="日志目录，默认是工具目录下的 logs")
    parser.add_argument("--headless", action="store_true", help="无界面模式，仅用于调试")
    return parser


def main() -> None:
    configure_console()
    parser = build_parser()
    args = parser.parse_args()

    interactive = not args.input
    if interactive:
        print("Amazon 关键词 ASIN 广告位置检测工具")
        print("只读取源 Excel；不会生成或修改任何 Excel 文件。检测结果实时显示并保存为日志。\n")
        source_default = str(DEFAULT_SOURCE)
        args.input = prompt_value("Excel 文件", source_default).strip('"')
        args.amazon = prompt_value("Amazon 站点", args.amazon)
        args.delay = float(prompt_value("关键词之间等待秒数", str(args.delay)))
        args.limit = int(prompt_value("本次扫描前几个唯一关键词（0 表示全部）", "3"))

    try:
        text_log, _ = run_tool(args)
        if interactive:
            answer = input("\n是否现在打开可读日志？[Y/n]：").strip().lower()
            if answer in {"", "y", "yes"}:
                os.startfile(text_log)  # type: ignore[attr-defined]
            input("检查完成后按 Enter 关闭窗口……")
    except Exception as error:
        print(f"\n错误：{error}", file=sys.stderr)
        if interactive:
            input("按 Enter 关闭窗口……")
        raise SystemExit(1) from error


if __name__ == "__main__":
    main()
